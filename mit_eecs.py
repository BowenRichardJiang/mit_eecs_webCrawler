import requests
from bs4 import BeautifulSoup
import os
from docx import Document
from docx.shared import Pt
from docx.shared import Inches

headers = {'User-Agent':"Mozilla/5.0 (Windows NT 6.1; WOW64) AppleWebKit/537.1 (KHTML, like Gecko) Chrome/22.0.1207.1 Safari/537.1"}
mit_url = 'http://www.eecs.mit.edu/'
mit_html = requests.get(mit_url,  headers=headers)
Soup=BeautifulSoup(mit_html.text,'lxml')
all_a=Soup.find('div',class_='item-list').find_all('a')

title =[]
for i in all_a:
    title.append(i.get_text())
num1 = 0

all_des=Soup.find_all('div',class_='field-content')
des = []
num2 = 0
for j in all_des:
    des.append(j.get_text().strip())

doc=Document()
prev = ''
for k in all_a:
    href = k['href']
    now = href
    if now == prev:
        continue
    prev = now
    html_url = requests.get(prev,headers=headers)
    html_Soup = BeautifulSoup(html_url.text, 'lxml')
    image_url = html_Soup.find('img')['src']
    image=requests.get(image_url,headers=headers)
    f=open('1.jpg','wb')
    f.write(image.content)
    f.close()
    num1 +=1
    num2 +=1
    if title[num1] == '' :
        for i in range(1,10):
            num1+=1
            if title[num1] != '':
                break
    if des[num2] == '':
        for i in range(1,10):
            num2+=1
            if des[num2] != '':
                break
    doc.add_picture('1.jpg',width=Inches(1.25))
    doc.add_paragraph(title[num1])
    doc.add_paragraph(des[num2])
doc.save('mit_eecs.docx')
os.remove('1.jpg')
